"""Carga de 22/10. Credenciais são lidas localmente e nunca impressas.
prepare: extrai e reconcilia. apply: backup + substituição transacional autorizada.
"""
import argparse, collections, datetime as dt, decimal, difflib, hashlib, json, os, re, sys
from fractions import Fraction
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path('C:/Users/lucas/Downloads/CENTROS_DE_CUSTO_CORRIGIDA - 22.10.xlsx')
DATA=ROOT/'admin/_financeiro_22_10.json'
D=decimal.Decimal
def text(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return re.sub(r'\s+',' ',str(v)).strip()
def money(v):
    if v is None or text(v)=='':return '0.00'
    return str(D(str(v)).quantize(D('.01'),rounding=decimal.ROUND_HALF_UP))
def php(v):
    if v is None:return 'null'
    if isinstance(v,(int,float)):return str(v)
    if isinstance(v,list):return '['+','.join(php(x) for x in v)+']'
    if isinstance(v,dict):return '['+','.join(php(k)+' => '+php(x) for k,x in v.items())+']'
    return "'"+str(v).replace('\\','\\\\').replace("'","\\'")+"'"
def extract():
    w=openpyxl.load_workbook(SOURCE,read_only=True,data_only=True)
    out=dict(source=SOURCE.name,sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(),lancamentos=[],jato=[],fixos=[],obras=[],funcionarios=[],cadastros=[],periodos=[],competencias={},pendencias=[])
    def issue(aba,linha,motivo,original=None):out['pendencias'].append(dict(aba=aba,linha=linha,motivo=motivo,original=original))
    def date(v,aba,n,col):
        if v is None or v=='':return None
        if isinstance(v,(dt.datetime,dt.date)):return v.strftime('%Y-%m-%d')
        if isinstance(v,(int,float)) and 30000<v<80000:return (dt.datetime(1899,12,30)+dt.timedelta(days=int(v))).strftime('%Y-%m-%d')
        s=text(v)
        for fmt in ['%d/%m/%Y','%Y-%m-%d']:
            try:return dt.datetime.strptime(s,fmt).strftime('%Y-%m-%d')
            except ValueError:pass
        issue(aba,n,'Data inválida em '+col+'; campo deixado vazio e original preservado.',s);return None
    def clean(v):return '' if text(v)=='-' else text(v)
    def installment(v):
        if isinstance(v,(int,float)) and 0<v<1:
            f=Fraction(v).limit_denominator(99);return str(f.numerator)+'/'+str(f.denominator)
        return text(v)
    cash_total=D('0'); source_rows=0
    for n,r in enumerate(w['Fluxo de Caixa'].iter_rows(min_row=2,values_only=True),2):
        r=list(r)+[None]*14
        if r[8] is None and r[9] is None:continue
        source_rows+=1;i=D(money(r[8]));j=D(money(r[9]));cash_total+=i+j
        ent=max(i,D(0))+max(j,D(0));sai=max(-i,D(0))+max(-j,D(0));kind='ENTRADA' if ent>0 and sai==0 else 'SAIDA'
        notes=[text(r[12])];issued=date(r[0],'Fluxo de Caixa',n,'Emissão');due=date(r[7],'Fluxo de Caixa',n,'Vencimento')
        for value,parsed,label in [(r[0],issued,'Emissão'),(r[7],due,'Vencimento')]:
            if value and not parsed:notes.append(label+' original: '+text(value))
        declared=text(r[1]).replace('Í','I')
        if declared!=kind or j>0 or i<0:
            issue('Fluxo de Caixa',n,'Tipo divergente; sinal original preservado no cálculo.',dict(tipo=text(r[1]),entrada=money(r[8]),saida=money(r[9])))
            notes.append('Tipo original: '+text(r[1])+'; classificação pelo sinal.')
        row=[kind,text(r[2]),re.sub(r'^OS\s*','',clean(r[3]),flags=re.I),clean(r[4]),text(r[5]),installment(r[6]),issued,due,str(ent.quantize(D('.01'))),str(sai.quantize(D('.01'))),text(r[11]),'GERAL',' | '.join(x for x in notes if x)]
        if len(row[12])>255:issue('Fluxo de Caixa',n,'Observação completa preservada nesta pendência.',row[12]);row[12]=row[12][:255]
        out['lancamentos'].append(row)
    for n,r in enumerate(w['TCM '].iter_rows(values_only=True),1):
        r=list(r)+[None]*15
        if not r[5] or r[12] is None:continue
        row=['SAIDA',text(r[6]),clean(r[7]),clean(r[8]),text(r[9]),installment(r[10]),date(r[4],'TCM',n,'Emissão'),date(r[11],'TCM',n,'Vencimento'),'0.00',str(abs(D(money(r[12])))), '', 'TCM','']
        matches=[x for x in out['lancamentos'] if x[:10]==row[:10]]
        if matches:matches[0][11]='TCM'
        else:issue('TCM',n,'Item não conciliado com o fluxo; confirmar se é rateio antes de lançar novamente.',row)
    for n,r in enumerate(w['JATO E PINTURA '].iter_rows(min_row=2,values_only=True),2):
        if r[3] is not None:out['jato'].append([text(r[1]),text(r[2]),date(r[3],'JATO E PINTURA',n,'Saída'),money(r[4]),money(r[5]),money(r[6])])
    salary_rows=list(w['Custos Fixos'].iter_rows(values_only=True))
    for r in salary_rows[1:8]:out['fixos'].append([text(r[0]),money(r[1])])
    for r in w['Custo por Obra'].iter_rows(min_row=2,values_only=True):
        if r[0]:out['obras'].append([re.sub(r'^OS\s*','',text(r[2])),text(r[0])])
    people={};period=None;closed=False
    for n,r in enumerate(salary_rows,1):
        if n<11:continue
        name=text(r[0]);value=r[1];vt=r[2];pix=text(r[3]);owner=text(r[4]);note=text(r[5])
        if name.startswith('SALARIO '):
            day=date(value,'Custos Fixos',n,'Período');kind='Adiantamento' if 'ADIANTAMENTO' in name else 'Fechamento'
            if not day:raise ValueError('Período sem data na linha '+str(n))
            period=[day,kind,[]];out['periodos'].append(period);closed=False
            d=dt.date.fromisoformat(day);competence=(d.replace(day=1)-dt.timedelta(days=1)).strftime('%Y-%m') if kind=='Fechamento' and d.day<=10 else d.strftime('%Y-%m')
            out['competencias'][day]=competence
            if kind=='Fechamento' and d.day>10:issue('Custos Fixos',n,'Cabeçalho informa Fechamento em data posterior ao dia 10; competência atribuída ao mês da data e deve ser conferida.',day)
            continue
        if not name and isinstance(value,(int,float)) and period:closed=True;continue
        if not name:continue
        if not isinstance(value,(int,float)) and not isinstance(vt,(int,float)):
            if value:issue('Custos Fixos',n,'Anotação sem valor de pagamento estruturado.',dict(nome=name,valor=text(value)))
            continue
        p=people.setdefault(name,dict(nome=name,pix=set(),banks=set(),notes=[]))
        if pix:
            if pix.upper().startswith('AG:'):p['banks'].add(pix)
            else:p['pix'].add(pix)
        if owner:p['notes'].append('Titular: '+owner)
        payment=[name,money(value),money(vt),' | '.join(x for x in [note,'Titular informado: '+owner if owner else ''] if x)]
        if not period or closed:issue('Custos Fixos',n,'Pagamento sem período próprio identificado; não atribuído a uma data por suposição.',payment);continue
        period[2].append(payment)
    for name,p in people.items():
        keys=sorted(p['pix']);key=keys[0] if len(keys)==1 else '';bank=' | '.join(sorted(p['banks']))
        if len(keys)>1:issue('Custos Fixos',0,'Chaves PIX divergentes: conferir cadastro de '+name,keys)
        if len(bank)>150:issue('Custos Fixos',0,'Dados bancários completos de '+name,bank);bank=bank[:150]
        cpf=key if re.fullmatch(r'\d{3}\.\d{3}\.\d{3}-\d{2}',key) and not p['notes'] else ''
        out['funcionarios'].append([name,key,bank]);out['cadastros'].append([name,cpf,'',key,bank])
    names=list(people)
    for i,a in enumerate(names):
        for b in names[i+1:]:
            if difflib.SequenceMatcher(None,a,b).ratio()>.87:issue('Custos Fixos',0,'Nomes semelhantes mantidos separados até conferência.',[a,b])
    entries=sum(D(r[8]) for r in out['lancamentos']);expenses=sum(D(r[9]) for r in out['lancamentos'])
    assert entries-expenses==cash_total and len(out['lancamentos'])==source_rows
    out['resumo']=dict(lancamentos=source_rows,entradas=str(entries),saidas=str(expenses),resultado=str(entries-expenses),funcionarios=len(out['cadastros']),periodos=len(out['periodos']),pagamentos=sum(len(p[2]) for p in out['periodos']),pendencias=len(out['pendencias']),jato=len(out['jato']),obras=len(out['obras']),custos_fixos=len(out['fixos']))
    DATA.write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
    (ROOT/'admin/_financeiro_dados.php').write_text('<?php\n// Carga reconciliada da planilha de 22/10/2026.\nreturn '+php(out)+';\n',encoding='utf-8')
    print(json.dumps(out['resumo'],ensure_ascii=True))
    return out

def connect():
    sys.path.insert(0,os.path.join(os.environ['TEMP'],'axion-finance-deps'));import pymysql
    source=(ROOT/'admin/_secrets.php').read_text(encoding='utf-8-sig');cfg={}
    for key in ['DB_HOST','DB_NAME','DB_USER','DB_PASS']:
        m=re.search(r"define\(\s*['\"]"+key+r"['\"]\s*,\s*(['\"])(.*?)\1\s*\)",source,re.S)
        if not m:raise RuntimeError('Configuração incompleta.')
        cfg[key]=m.group(2).replace("\\'","'").replace('\\\\','\\')
    return pymysql.connect(host=cfg['DB_HOST'],user=cfg['DB_USER'],password=cfg['DB_PASS'],database=cfg['DB_NAME'],charset='utf8mb4',autocommit=False,connect_timeout=10,cursorclass=pymysql.cursors.DictCursor)

def apply():
    data=json.loads(DATA.read_text(encoding='utf-8'));assert data['sha256']==hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    db=connect();tables=['centro_folha_pagamentos','centro_folha_periodos','centro_funcionarios','centro_custos_obra','centro_obras','centro_lancamentos','centro_jato_pintura','centro_custos_fixos']
    backup_dir=ROOT/'admin/_financeiro_backups';backup_dir.mkdir(exist_ok=True);(backup_dir/'.htaccess').write_text('Require all denied\n')
    backup=backup_dir/('antes-22-10-'+dt.datetime.now().strftime('%Y%m%d-%H%M%S')+'.json')
    try:
        with db.cursor() as c:
            c.execute("SELECT GET_LOCK('axion_financeiro_inicial',15) AS acquired");assert c.fetchone()['acquired']==1
            c.execute('SELECT id FROM centro_audit_log WHERE detalhes LIKE %s LIMIT 1',('%'+data['sha256']+'%',))
            if c.fetchone():raise RuntimeError('Esta versão da planilha já foi aplicada. Reimportação bloqueada para preservar alterações feitas no site.')
            c.execute("SELECT TABLE_NAME,ENGINE FROM information_schema.TABLES WHERE TABLE_SCHEMA=DATABASE()");engines={r['TABLE_NAME']:r['ENGINE'] for r in c.fetchall()}
            assert all(engines.get(t)=='InnoDB' for t in tables+['centro_audit_log'])
            preserved={}
            for t in ['users','propostas','producao_os','producao_pecas']:
                if t in engines:c.execute('SELECT COUNT(*) AS n FROM '+t);preserved[t]=c.fetchone()['n']
            c.execute('SELECT TABLE_NAME,REFERENCED_TABLE_NAME FROM information_schema.KEY_COLUMN_USAGE WHERE REFERENCED_TABLE_SCHEMA=DATABASE() AND REFERENCED_TABLE_NAME IS NOT NULL')
            for relation in c.fetchall():
                if relation['REFERENCED_TABLE_NAME'] in tables and relation['TABLE_NAME'] not in tables:raise RuntimeError('Há dependência externa ao financeiro; não foi apagado nenhum dado.')
            snap={};schemas={}
            for t in tables:
                c.execute('SELECT * FROM '+t);snap[t]=c.fetchall();c.execute('SHOW CREATE TABLE '+t);schemas[t]=c.fetchone()['Create Table']
            backup.write_text(json.dumps(dict(source=data['source'],tabelas=snap,schemas=schemas),ensure_ascii=False,default=str,indent=2),encoding='utf-8')
            assert json.loads(backup.read_text(encoding='utf-8'))['tabelas'].keys()==snap.keys()
            db.commit()
            c.execute('SHOW COLUMNS FROM centro_folha_periodos');cols={r['Field'] for r in c.fetchall()}
            if 'competencia' not in cols:c.execute('ALTER TABLE centro_folha_periodos ADD competencia CHAR(7) NULL')
            c.execute('SHOW COLUMNS FROM centro_folha_pagamentos');cols={r['Field'] for r in c.fetchall()}
            for col,definition in [('situacao_pagamento',"VARCHAR(20) NOT NULL DEFAULT 'a_confirmar'"),('valor_pago','DECIMAL(12,2) NULL'),('data_pagamento','DATE NULL')]:
                if col not in cols:c.execute('ALTER TABLE centro_folha_pagamentos ADD '+col+' '+definition)
            c.execute("CREATE TABLE IF NOT EXISTS centro_import_pendencias (id INT UNSIGNED AUTO_INCREMENT PRIMARY KEY, origem VARCHAR(255) NOT NULL, sha256 CHAR(64) NOT NULL, dados LONGTEXT NOT NULL, resolvida TINYINT NOT NULL DEFAULT 0, created_at DATETIME DEFAULT CURRENT_TIMESTAMP) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4")
            db.commit();db.begin()
            # DELETE mantém AUTO_INCREMENT e evita reutilizar IDs referidos pela auditoria antiga.
            for t in tables:c.execute('DELETE FROM '+t)
            groups=[('centro_lancamentos','tipo,descricao,os,nfe,cliente_fornecedor,parcela,data_lancamento,vencimento,valor_entrada,valor_saida,situacao,centro,observacao','lancamentos'),('centro_jato_pintura','cliente,produto,data_saida,kg,vr_kg,vr_total','jato'),('centro_custos_fixos','descricao,valor','fixos'),('centro_obras','os,nome','obras'),('centro_funcionarios','nome,cpf,telefone,pix_chave,banco_info','cadastros')]
            for table,columns,key in groups:
                rows=data[key]
                if rows:c.executemany('INSERT INTO '+table+'('+columns+') VALUES('+','.join(['%s']*len(rows[0]))+')',rows)
            c.execute('SELECT id,nome FROM centro_funcionarios');ids={r['nome']:r['id'] for r in c.fetchall()}
            for date,kind,payments in data['periodos']:
                c.execute('INSERT INTO centro_folha_periodos(data_periodo,tipo,competencia) VALUES(%s,%s,%s)',(date,kind,data['competencias'][date]));period=c.lastrowid
                for name,value,vt,note in payments:c.execute('INSERT INTO centro_folha_pagamentos(periodo_id,funcionario_id,valor,vale_transporte,observacao,situacao_pagamento) VALUES(%s,%s,%s,%s,%s,%s)',(period,ids[name],value,vt,note,'a_confirmar'))
            for item in data['pendencias']:c.execute('INSERT INTO centro_import_pendencias(origem,sha256,dados) VALUES(%s,%s,%s)',(data['source'],data['sha256'],json.dumps(item,ensure_ascii=False)))
            c.execute('SELECT COUNT(*) AS n,SUM(valor_entrada) AS e,SUM(valor_saida) AS s FROM centro_lancamentos');actual=c.fetchone()
            assert actual['n']==data['resumo']['lancamentos'] and actual['e']==D(data['resumo']['entradas']) and actual['s']==D(data['resumo']['saidas'])
            for table,key in [('centro_funcionarios','funcionarios'),('centro_folha_pagamentos','pagamentos'),('centro_folha_periodos','periodos')]:
                c.execute('SELECT COUNT(*) AS n FROM '+table);assert c.fetchone()['n']==data['resumo'][key]
            for t in ['users','propostas','producao_os','producao_pecas']:
                if t in engines:
                    c.execute('SELECT COUNT(*) AS n FROM '+t);after=c.fetchone()['n'];assert after==preserved[t];print('preservado',t,after)
            details=dict(atualizacao_excel=data['sha256'],origem=data['source'],contagens=data['resumo'],backup=backup.name,competencia='Fechamentos até dia 10: mês anterior; demais: mês da data. Conferir cabeçalhos divergentes.')
            c.execute('INSERT INTO centro_audit_log(user_id,username,acao,tabela,registro_id,detalhes,ip_hash) VALUES(NULL,%s,%s,%s,0,%s,%s)',('carga_planilha_cli','create','centro_lancamentos',json.dumps(details,ensure_ascii=False),hashlib.sha256(b'authorized-excel-refresh').hexdigest()))
            db.commit();print('COMMIT VERIFIED',json.dumps(data['resumo']));print('Backup local:',str(backup))
    except Exception:
        db.rollback();raise
    finally:
        with db.cursor() as c:c.execute("SELECT RELEASE_LOCK('axion_financeiro_inicial')")
        db.close()

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('action',choices=['prepare','apply']);args=parser.parse_args()
    if args.action=='prepare':extract()
    else:apply()
